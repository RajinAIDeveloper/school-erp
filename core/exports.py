import csv
from io import BytesIO
from xml.sax.saxutils import escape

from django.http import HttpResponse


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def spreadsheet(title, headers, rows, fmt="csv", extra_sheets=()):
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        for index, (name, cols, data) in enumerate([(title, headers, rows), *extra_sheets]):
            ws = wb.active if index == 0 else wb.create_sheet()
            ws.title = name[:31]
            ws.append(cols)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for row in data:
                ws.append([safe_cell(v) for v in row])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(45, max(12, max(len(str(c.value or "")) for c in col) + 2))
        output = BytesIO()
        wb.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response.write("\ufeff")
        writer = csv.writer(response)
        writer.writerow(headers)
        writer.writerows([[safe_cell(v) for v in row] for row in rows])
        fmt = "csv"
    response["Content-Disposition"] = f'attachment; filename="{title}.{fmt}"'
    return response


def pdf_response(title, headers, rows, subtitle=""):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    output = BytesIO()
    styles = getSampleStyleSheet()
    page = landscape(A4) if len(headers) > 6 else A4
    doc = SimpleDocTemplate(output, pagesize=page, rightMargin=25, leftMargin=25)
    data = [[Paragraph(escape(str(v if v is not None else "")), styles["Normal"]) for v in row] for row in [headers, *list(rows)]]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#e2e8f0")), ("GRID", (0,0),(-1,-1),0.4,colors.grey), ("VALIGN",(0,0),(-1,-1),"TOP"), ("BOTTOMPADDING",(0,0),(-1,-1),7)]))
    doc.build([Paragraph(escape(title), styles["Title"]), Paragraph(escape(subtitle), styles["Normal"]), Spacer(1,15), table])
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="report.pdf"'
    return response
