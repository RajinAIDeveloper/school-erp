"""
Class result exports: grade distribution, tabulation sheet, merit lists, and what every
export says about itself. Figures come from the same rows as the cards, so they reconcile.
"""

import csv
import io
from datetime import date
from decimal import Decimal

from django.test import Client
from openpyxl import load_workbook

from academics.models import ClassLevel, Section, Subject
from core.models import AssessmentSystem
from examinations.exports import grade_distribution, merit_list, tabulation
from examinations.models import Exam, ExamSchedule, ResultSnapshot
from examinations.presets import install_preset
from examinations.rulebooks import RULEBOOKS
from examinations.services import build_result_sheet, publish_exam, save_mark
from students.models import Enrollment, Student
from tests.test_board_results import mark, mark_everything


def login(user):
    client = Client()
    client.force_login(user)
    return client


def results_url(exam, level, **extra):
    query = "&".join(f"{k}={v}" for k, v in {"exam": exam.pk, "class_level": level.pk, **extra}.items())
    return f"/exams/results/?{query}"


def csv_rows(response):
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


# ============================================================ grade distribution


def test_each_subject_is_counted_over_the_students_who_sat_it(erp, board):
    mark_everything(erp, board)
    sheet = build_result_sheet(board.exam, board.level)
    letters, lines = grade_distribution(sheet, sheet["rows"])
    by_subject = {line["subject"]: line for line in lines}
    # Physics is Science only, Geography Humanities only: one student each, not two.
    assert by_subject["Physics"]["sat"] == 1
    assert by_subject["Geography and Environment"]["sat"] == 1
    assert by_subject["General Math"]["sat"] == 2
    # 85 on every paper is an A+ on the national scale, so 100% at or above A+.
    assert by_subject["General Math"]["counts"]["A+"] == 2
    assert by_subject["General Math"]["at_or_above"]["A+"] == 100.0
    assert letters[0] == "A+" and letters[-1] == "F"


def test_an_absence_is_counted_apart_and_not_in_the_percentages(erp, board):
    for code, schedule in board.schedules.items():
        for enrollment in (board.science, board.humanities):
            from examinations.subjects import takes_paper

            if takes_paper(enrollment, schedule):
                absent = code == "109" and enrollment == board.humanities
                mark(erp, schedule, enrollment, 0 if absent else 85, absent=absent)
    sheet = build_result_sheet(board.exam, board.level)
    _letters, lines = grade_distribution(sheet, sheet["rows"])
    maths = next(line for line in lines if line["subject"] == "General Math")
    assert (maths["sat"], maths["absent"]) == (1, 1)
    assert maths["at_or_above"]["A+"] == 100.0


def test_the_distribution_screen_and_downloads(erp, board):
    mark_everything(erp, board)
    client = login(erp.admin)
    page = client.get(results_url(board.exam, board.level, report="distribution"))
    assert page.status_code == 200
    assert b"<th>Sat</th>" in page.content and b"<th>A+</th>" in page.content
    download = csv_rows(client.get(results_url(board.exam, board.level, report="distribution", format="csv")))
    assert download[0] == ["Exam", f"Half Yearly ({erp.year})"]
    assert ["Version", "Draft: not published, and may still change"] in download
    workbook = load_workbook(
        io.BytesIO(client.get(results_url(board.exam, board.level, report="distribution", format="xlsx")).content)
    )
    assert workbook.sheetnames == ["grade-distribution", "At or above (%)", "About"]
    pdf = client.get(results_url(board.exam, board.level, report="distribution", format="pdf"))
    assert pdf["Content-Type"] == "application/pdf"


# ============================================================ tabulation


def test_the_tabulation_sheet_matches_the_published_cards(erp, board):
    mark_everything(erp, board, {(board.science.pk, "126"): 60})
    publish_exam(board.exam, erp.admin)
    sheet = build_result_sheet(board.exam, board.level)
    headers, body = tabulation(sheet["rows"])
    assert headers[:4] == ["Roll", "Student", "Group", "4th subject"]
    assert "Bangla marks" in headers and "Bangla 1st paper marks" not in headers
    for snapshot in ResultSnapshot.objects.filter(exam=board.exam):
        payload = snapshot.payload
        line = next(r for r in body if r[1] == payload["student"])
        assert line[headers.index("GPA")] == payload["gpa"]
        assert line[headers.index("Result")] == payload["result"]
        bangla = next(u for u in payload["subjects"] if u["name"] == "Bangla")
        assert line[headers.index("Bangla marks")] == bangla["score"]
    science = next(r for r in body if r[1] == "Nabila")
    assert science[headers.index("4th subject")] == "Higher Math"
    # A subject a student does not take is blank, not a zero or a dash.
    assert science[headers.index("Geography and Environment marks")] == ""


def test_the_tabulation_sheet_is_only_for_national_classes(erp):
    exam, level = cambridge_class(erp)
    page = login(erp.admin).get(results_url(exam, level, report="tabulation"))
    assert b"tabulation sheet is for classes following the national curriculum" in page.content


# ============================================================ merit lists


def test_a_merit_list_by_group_ranks_only_that_group(erp, board):
    mark_everything(erp, board, {(board.humanities.pk, "109"): 95})
    client = login(erp.admin)
    rows = csv_rows(client.get(results_url(board.exam, board.level, report="merit", group="Science", format="csv")))
    assert ["Filter", "Science"] in rows
    header = rows.index(next(r for r in rows if r and r[0] == "Position"))
    listed = rows[header + 1 :]
    assert [r[1] for r in listed] == ["Nabila"]
    assert listed[0][0] == "1"


def test_students_without_a_complete_result_are_listed_but_not_ranked(erp, board):
    mark_everything(erp, board)
    from examinations.models import Mark

    Mark.objects.filter(enrollment=board.humanities, schedule=board.schedules["109"]).delete()
    sheet = build_result_sheet(board.exam, board.level)
    ranked, unranked = merit_list(sheet["rows"])
    assert [r["student"] for r in ranked] == ["Nabila"]
    assert [r["student"] for r in unranked] == ["Rupa"]
    page = login(erp.admin).get(results_url(board.exam, board.level, report="merit"))
    assert b"Not ranked" in page.content
    assert b"1 without a complete result are listed after, not ranked" in page.content


def cambridge_class(erp):
    scale, _ = install_preset(erp.school, "cambridge-igcse")
    level = ClassLevel.objects.create(
        school=erp.school, name="Year 11", order=11, assessment_system=AssessmentSystem.CAMBRIDGE
    )
    section = Section.objects.create(school=erp.school, class_level=level, name="Green", shift="morning")
    exam = Exam.objects.create(school=erp.school, academic_year=erp.year, name="Mock", grade_scale=scale)
    subject = Subject.objects.create(school=erp.school, name="Chemistry", code="0620")
    paper = ExamSchedule.objects.create(
        school=erp.school, exam=exam, class_level=level, subject=subject, full_marks=100, pass_marks=0
    )
    for roll, (name, score) in enumerate((("Ayan", 91), ("Mira", 72)), 1):
        pupil = Enrollment.objects.create(
            school=erp.school,
            student=Student.objects.create(
                school=erp.school,
                student_id=f"Y11-{roll}",
                first_name=name,
                gender="F",
                date_of_birth=date(2009, 1, 1),
                admission_date=date(2026, 1, 1),
            ),
            academic_year=erp.year,
            class_level=level,
            section=section,
            roll_number=roll,
        )
        save_mark(user=erp.admin, schedule=paper, enrollment=pupil, score=Decimal(score))
    return exam, level


def test_no_merit_list_where_the_rulebook_has_no_positions_until_the_exam_turns_them_on(erp):
    exam, level = cambridge_class(erp)
    client = login(erp.admin)
    page = client.get(results_url(exam, level, report="merit"))
    assert b"Positions are off for this exam" in page.content
    exam.show_rank = True
    exam.save()
    rows = csv_rows(client.get(results_url(exam, level, report="merit", format="csv")))
    header = next(r for r in rows if r and r[0] == "Position")
    assert "GPA" not in header and "Result" not in header
    assert ["1", "Ayan"] == rows[rows.index(header) + 1][:2]


def test_the_cambridge_distribution_uses_cambridge_letters(erp):
    exam, level = cambridge_class(erp)
    sheet = build_result_sheet(exam, level)
    letters, lines = grade_distribution(sheet, sheet["rows"])
    assert letters[:3] == ["A*", "A", "B"]
    assert lines[0]["counts"]["A*"] == 1 and lines[0]["counts"]["B"] == 1
    assert lines[0]["at_or_above"]["A"] == 50.0


# ============================================================ the class sheet


def test_the_class_sheet_says_which_version_it_is(erp, board):
    mark_everything(erp, board)
    client = login(erp.admin)
    draft = csv_rows(client.get(results_url(board.exam, board.level, format="csv")))
    assert ["Version", "Draft: not published, and may still change"] in draft
    publish_exam(board.exam, erp.admin)
    published = csv_rows(client.get(results_url(board.exam, board.level, format="csv")))
    assert ["Version", "Published version 1"] in published
    assert ["Rulebook", RULEBOOKS["national"].label] in published
    workbook = load_workbook(io.BytesIO(client.get(results_url(board.exam, board.level, format="xlsx")).content))
    assert workbook.sheetnames[0] == "results" and "About" in workbook.sheetnames


def test_a_filter_keeps_the_summary_to_the_students_shown(erp, board):
    mark_everything(erp, board)
    page = login(erp.admin).get(results_url(board.exam, board.level, group="Humanities"))
    assert b"Complete 1" in page.content
    assert b"Rupa" in page.content and b"Nabila" not in page.content
