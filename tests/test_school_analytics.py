"""
Phase 6d: the whole school in one published exam, for the Principal, Vice Principal and
Administrator. Class against class, every subject in every class, groups and boys and girls,
the top of each class, who has improved and who has slipped, and the year's trend. Nobody
else opens it, and no school sees another's.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import Client
from openpyxl import load_workbook

from analytics.school import school_overview, year_trend
from examinations.models import Exam, ExamSchedule
from examinations.services import publish_exam, save_mark
from tests.test_analytics_foundation import add_user
from tests.test_progress import two_exams


def login(user):
    client = Client()
    client.force_login(user)
    return client


def test_the_school_figures_come_from_the_published_results(erp, board):
    two_exams(erp, board)
    data = school_overview(board.exam)
    assert data["students"] == 2 and data["failing"] == 0
    assert data["average"] == Decimal("80.2") and data["gpa"] == Decimal("5.0")
    assert data["pass_rate"] == Decimal("100.0")
    (level,) = data["classes"]
    assert level["level"] == "Class 9" and level["highest"] == Decimal("85.8")
    assert [s["name"] for s in level["sections"]] == [board.section.name]
    assert [f.student.first_name for f in level["top"]] == ["Nabila", "Rupa"]
    # Every subject in every class: Bangla's two papers count once, as the subject.
    cells = dict(data["grid"][0]["cells"])
    bangla = dict(data["subjects"])
    bangla_pk = next(pk for pk, name in bangla.items() if name == "Bangla")
    assert cells[bangla_pk]["average"] == Decimal("77.5") and cells[bangla_pk]["students"] == 2
    assert {g["name"]: g["average"] for g in data["groups"]} == {
        "Humanities": Decimal("74.5"),
        "Science": Decimal("85.8"),
    }


def test_boys_and_girls_are_compared(erp, board):
    two_exams(erp, board)
    rupa = board.humanities.student
    rupa.gender = "M"
    rupa.save()
    genders = {g["name"]: (g["students"], g["average"]) for g in school_overview(board.exam)["genders"]}
    assert genders == {"Boys": (1, Decimal("74.5")), "Girls": (1, Decimal("85.8"))}


def test_the_movers_and_the_trend_follow_the_years_exams(erp, board):
    two_exams(erp, board)
    data = school_overview(board.exam)
    assert [(m["fact"].student.first_name, m["change"]) for m in data["improved"]] == [("Nabila", Decimal("4.2"))]
    assert [(m["fact"].student.first_name, m["change"]) for m in data["slipped"]] == [("Rupa", Decimal("-4.2"))]
    trend = data["trend"]
    assert trend["labels"] == ["First Term", "Half Yearly"]
    # One class sat both exams: its own line, and no separate school line to repeat it.
    assert trend["series"] == [("Class 9", [Decimal("80.2"), Decimal("80.2")])]
    # The first exam of the year has nothing earlier to compare with.
    first = Exam.objects.get(school=erp.school, name="First Term")
    assert school_overview(first)["improved"] == [] and school_overview(first)["slipped"] == []


def two_classes_two_exams(erp):
    """Class 1 and Class 2 both sit Term 1 and Term 2, one student each, in Math."""
    from academics.models import ClassLevel, Section
    from students.models import Enrollment, Student

    level = ClassLevel.objects.create(school=erp.school, name="Class 2", order=2)
    section = Section.objects.create(school=erp.school, class_level=level, name="A")
    student = Student.objects.create(
        school=erp.school,
        student_id="S2",
        first_name="Tanvir",
        gender="M",
        date_of_birth=date(2015, 1, 1),
        admission_date=date(2026, 1, 1),
    )
    older = Enrollment.objects.create(
        school=erp.school, student=student, academic_year=erp.year, class_level=level, section=section, roll_number=1
    )
    erp.exam.end_date = date(2026, 4, 1)
    erp.exam.save()
    later = Exam.objects.create(
        school=erp.school, academic_year=erp.year, name="Term 2", grade_scale=erp.scale, end_date=date(2026, 8, 1)
    )
    scores = {(erp.exam, erp.level): 60, (erp.exam, level): 80, (later, erp.level): 70, (later, level): 90}
    for exam in (erp.exam, later):
        for lvl, enrollment in ((erp.level, erp.enrollment), (level, older)):
            schedule, _ = ExamSchedule.objects.get_or_create(
                school=erp.school, exam=exam, class_level=lvl, subject=erp.subject, full_marks=100, pass_marks=33
            )
            save_mark(user=erp.admin, schedule=schedule, enrollment=enrollment, score=Decimal(scores[exam, lvl]))
        publish_exam(exam, erp.admin)
        exam.refresh_from_db()
    return later


def test_a_school_line_is_drawn_only_when_every_exam_had_the_same_classes(erp, board):
    later = two_classes_two_exams(erp)
    trend = dict(year_trend(later)["series"])
    assert trend == {
        "School": [Decimal("70.0"), Decimal("80.0")],
        "Class 1": [Decimal("60.0"), Decimal("70.0")],
        "Class 2": [Decimal("80.0"), Decimal("90.0")],
    }
    # Once Class 9 has exams of its own, the classes no longer sat the same exams.
    two_exams(erp, board)
    assert "School" not in dict(year_trend(later)["series"])


def test_only_the_schools_managers_open_it(erp, board):
    two_exams(erp, board)
    url = f"/analytics/school/?exam={board.exam.pk}"
    for user in (erp.admin, add_user(erp, "principal", "Principal"), add_user(erp, "vice", "Vice Principal")):
        page = login(user).get(url).content.decode()
        assert "Class by class" in page and "Nabila" in page and "Most improved" in page
    assert login(erp.teacher).get(url).status_code == 403
    assert login(erp.accountant).get(url).status_code == 403
    assert login(erp.parent).get(url).status_code == 403
    # The analytics home offers the page to managers only.
    assert "School analytics" in login(erp.admin).get("/analytics/").content.decode()
    assert "School analytics" not in login(erp.teacher).get("/analytics/").content.decode()


def test_another_schools_exam_and_an_unpublished_one_are_not_found(erp, board):
    two_exams(erp, board)
    theirs = Exam.objects.create(
        school=erp.other, academic_year=erp.year, name="Theirs", grade_scale=erp.scale, status="published"
    )
    client = login(erp.admin)
    assert client.get(f"/analytics/school/?exam={theirs.pk}").status_code == 404
    draft = Exam.objects.create(school=erp.school, academic_year=erp.year, name="Weekly test", grade_scale=erp.scale)
    assert client.get(f"/analytics/school/?exam={draft.pk}").status_code == 404
    # An id that is not a number finds nothing, rather than failing.
    assert client.get("/analytics/school/?exam=abc").status_code == 404
    assert client.get("/analytics/paper/?exam=abc&subject=1").status_code == 404
    assert client.get("/analytics/section/?exam=1&section=x").status_code == 404


def test_nothing_published_yet_says_so(erp):
    page = login(erp.admin).get("/analytics/school/").content.decode()
    assert "No exam has been published this year yet." in page


def test_the_excel_export_has_a_sheet_for_each_view(erp, board):
    two_exams(erp, board)
    response = login(erp.admin).get(f"/analytics/school/?exam={board.exam.pk}&format=xlsx")
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames[1:] == ["Subjects", "Most improved", "Biggest drops"]
    first = book.worksheets[0]
    assert first["A1"].value == "Class" and first["A2"].value == "Class 9" and first["B2"].value == 2
    subjects = book["Subjects"]
    assert "Bangla" in [c.value for c in subjects[1]]
    assert book["Most improved"]["A2"].value == "Nabila"


def test_a_cambridge_class_has_no_pass_rate(erp, igcse):
    for pupil, score in zip(igcse.pupils, (92, 34), strict=True):
        save_mark(
            user=erp.admin,
            schedule=igcse.papers["0625"],
            enrollment=pupil,
            components={"mcq": "30", "theory": "60", "practical": "30"},
        )
        save_mark(user=erp.admin, schedule=igcse.papers["0510"], enrollment=pupil, score=Decimal(score))
    publish_exam(igcse.exam, erp.admin)
    igcse.exam.refresh_from_db()
    data = school_overview(igcse.exam)
    assert data["pass_rate"] is None and data["classes"][0]["pass_rate"] is None
    assert all(cell is None or cell["pass_rate"] is None for _pk, cell in data["grid"][0]["cells"])
    page = login(erp.admin).get(f"/analytics/school/?exam={igcse.exam.pk}").content.decode()
    assert "Year 10" in page and "Zara" in page


def test_the_page_reads_in_bangla(erp, board):
    two_exams(erp, board)
    erp.admin.language = "bn"
    erp.admin.save()
    page = login(erp.admin).get(f"/analytics/school/?exam={board.exam.pk}").content.decode()
    assert "বিদ্যালয়ের বিশ্লেষণ" in page and "শ্রেণি অনুযায়ী" in page
