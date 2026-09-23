from datetime import date, time, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from academics.models import AcademicYear, ClassLevel, Section
from attendance.models import StudentAttendance
from attendance.services import monthly_matrix, save_register
from core.models import School
from examinations.models import GradeRule, Mark, ResultSnapshot, UnlockRequest
from examinations.services import build_result_sheet, publish_exam, review_unlock, save_mark
from fees.models import FeeInvoice, FeePayment
from fees.services import cancel_payment, collect_payment, generate_invoices
from finance.models import Account, JournalEntry, record_simple_entry
from holidays.models import Holiday
from messaging.backends import BaseSMSBackend
from messaging.models import SMSMessage
from messaging.services import queue_batch
from students.forms import EnrollmentForm
from students.models import Enrollment, Student
from students.services import import_students, promote
from timetable.models import Period, RoutineSlot


def pay(erp, invoice, amount=400):
    return collect_payment(
        school=erp.school,
        user=erp.accountant,
        invoice=invoice,
        amount=amount,
        method="cash",
        reference="",
        date=date(2026, 9, 21),
    )


def test_fee_generation_is_idempotent(erp, invoice):
    result = generate_invoices(
        school=erp.school,
        user=erp.admin,
        academic_year=erp.year,
        class_level=erp.level,
        month=9,
        issue_date=date(2026, 9, 1),
        due_date=date(2026, 9, 10),
    )
    assert result == (0, 1)
    assert FeeInvoice.objects.count() == 1


def test_partial_and_final_payment_ledger(erp, invoice):
    payment = pay(erp, invoice)
    invoice.refresh_from_db()
    assert invoice.balance == 600 and invoice.status == "partial"
    assert payment.journal_entry.is_balanced
    pay(erp, invoice, 600)
    invoice.refresh_from_db()
    assert invoice.balance == 0 and invoice.status == "paid"
    assert Account.objects.get(school=erp.school, code="1010").balance() == 1000


def test_overpayment_rolls_back(erp, invoice):
    with pytest.raises(ValidationError):
        pay(erp, invoice, 1001)
    assert not FeePayment.objects.exists()
    assert not JournalEntry.objects.exists()


def test_payment_reversal_nets_zero_and_is_idempotent(erp, invoice):
    payment = pay(erp, invoice, 1000)
    cancel_payment(school=erp.school, user=erp.accountant, payment=payment, reason="Duplicate receipt")
    cancel_payment(school=erp.school, user=erp.accountant, payment=payment, reason="Duplicate receipt")
    assert Account.objects.get(school=erp.school, code="1010").balance() == 0
    assert Account.objects.get(school=erp.school, code="4090").balance() == 0
    invoice.refresh_from_db()
    assert invoice.balance == 1000
    assert JournalEntry.objects.count() == 2


def test_cross_school_payment_rejected(erp, invoice):
    with pytest.raises(PermissionDenied):
        collect_payment(
            school=erp.other,
            user=erp.admin,
            invoice=invoice,
            amount=100,
            method="cash",
            reference="",
            date=date.today(),
        )


def test_cross_school_account_rejected(erp):
    external = Account.objects.create(school=erp.other, code="1", name="Other", account_type="asset")
    with pytest.raises(PermissionDenied):
        record_simple_entry(
            erp.school,
            date.today(),
            "Wrong school",
            external,
            Account.objects.get(school=erp.school, code="4090"),
            100,
            source="manual",
        )


def test_publish_requires_complete_marks(erp):
    with pytest.raises(ValidationError):
        publish_exam(erp.exam, erp.admin)
    assert not ResultSnapshot.objects.exists()


def test_mark_validation_and_conflict(erp):
    with pytest.raises(ValidationError):
        save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(101))
    mark = save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    with pytest.raises(ValidationError):
        save_mark(
            user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(90), expected_version=0
        )
    assert Mark.objects.get(pk=mark.pk).marks_obtained == 80


def test_published_snapshot_immutable_and_lock(erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    exam = publish_exam(erp.exam, erp.admin)
    snap = ResultSnapshot.objects.get()
    old = snap.payload
    GradeRule.objects.filter(scale=erp.scale).update(grade_point=0)
    assert build_result_sheet(exam, erp.level)["rows"][0]["gpa"] == "5.00"
    with pytest.raises(ValidationError):
        save_mark(
            user=erp.admin, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(90), expected_version=1
        )
    snap.refresh_from_db()
    assert snap.payload == old


def test_scoped_unlock_new_snapshot_keeps_history(erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(70))
    publish_exam(erp.exam, erp.admin)
    unlock = UnlockRequest.objects.create(
        school=erp.school, schedule=erp.schedule, requested_by=erp.teacher, reason="Correction"
    )
    review_unlock(unlock, erp.admin, True)
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80), expected_version=1)
    assert list(ResultSnapshot.objects.order_by("version").values_list("version", flat=True)) == [1, 2]
    assert ResultSnapshot.objects.get(version=1).payload["total"] == "70.00"
    assert ResultSnapshot.objects.get(version=2).payload["total"] == "80.00"


def test_expired_unlock_rejected(erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(70))
    publish_exam(erp.exam, erp.admin)
    UnlockRequest.objects.create(
        school=erp.school,
        schedule=erp.schedule,
        requested_by=erp.teacher,
        reason="Old",
        status="approved",
        expires_at=timezone.now() - timedelta(seconds=1),
    )
    with pytest.raises(ValidationError):
        save_mark(
            user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80), expected_version=1
        )


def test_incomplete_absent_and_ranks(erp):
    from examinations.services import assign_ranks

    rows = [
        {"total": "90", "complete": True},
        {"total": "90", "complete": True},
        {"total": "80", "complete": True},
        {"total": "99", "complete": False},
    ]
    assign_ranks(rows)
    assert [r["rank"] for r in rows] == [1, 1, 3, None]
    before = build_result_sheet(erp.exam, erp.level)
    assert before["rows"][0]["result"] == "INCOMPLETE"
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=None, absent=True)
    after = build_result_sheet(erp.exam, erp.level)
    assert after["rows"][0]["result"] == "FAIL"
    assert after["subject_stats"][0]["average"] is None
    assert after["subject_stats"][0]["absent"] == 1


def test_unassigned_teacher_cannot_mark(erp):
    erp.enrollment.section = erp.other_section
    erp.enrollment.save()
    with pytest.raises(PermissionDenied):
        save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))


def test_attendance_holiday_and_duplicate(erp):
    day = date(2026, 9, 21)
    data = dict(school=erp.school, user=erp.teacher, day=day, entries=[(erp.enrollment, "present", "", None, None)])
    save_register(**data)
    save_register(**data)
    assert StudentAttendance.objects.count() == 1
    Holiday.objects.create(school=erp.school, name="Closed", start_date=day, end_date=day)
    with pytest.raises(ValidationError):
        save_register(**data)


def test_attendance_roster_permission(erp):
    erp.enrollment.section = erp.other_section
    erp.enrollment.save()
    with pytest.raises(PermissionDenied):
        save_register(
            school=erp.school,
            user=erp.teacher,
            day=date(2026, 9, 21),
            entries=[(erp.enrollment, "present", "", None, None)],
        )
    assert StudentAttendance.objects.count() == 0


def test_attendance_report_excludes_holidays(erp):
    Holiday.objects.create(school=erp.school, name="Closed", start_date=date(2026, 9, 1), end_date=date(2026, 9, 30))
    _, rows = monthly_matrix(erp.school, [erp.enrollment], 2026, 9)
    assert rows[0]["working"] == 0 and rows[0]["pct"] is None


def test_enrollment_cannot_use_other_school(erp):
    foreign = ClassLevel.objects.create(school=erp.other, name="Other", order=1)
    form = EnrollmentForm(
        {
            "student": erp.student.pk,
            "academic_year": erp.year.pk,
            "class_level": foreign.pk,
            "section": erp.section.pk,
            "roll_number": 2,
            "status": "enrolled",
        },
        school=erp.school,
    )
    assert not form.is_valid() and "class_level" in form.errors


def test_duplicate_enrollment_is_form_error(erp):
    form = EnrollmentForm(
        {
            "student": erp.student.pk,
            "academic_year": erp.year.pk,
            "class_level": erp.level.pk,
            "section": erp.section.pk,
            "roll_number": 1,
            "status": "enrolled",
        },
        school=erp.school,
    )
    assert not form.is_valid()


def test_promotion_keeps_history(erp):
    year = AcademicYear.objects.create(
        school=erp.school, name="2027", start_date=date(2027, 1, 1), end_date=date(2027, 12, 31)
    )
    n = promote(erp.school, erp.year, erp.section, year, erp.other_section)
    assert n == 1 and Enrollment.objects.count() == 2
    erp.enrollment.refresh_from_db()
    assert erp.enrollment.status == "promoted"


def test_invalid_import_is_atomic(erp):
    upload = SimpleUploadedFile(
        "students.csv",
        b"student_id,first_name,gender,date_of_birth,admission_date,status\nS2,Valid,M,2016-01-01,2026-01-01,active\nS3,Invalid,X,wrong,2026-01-01,active",
    )
    with pytest.raises(ValidationError):
        import_students(erp.school, upload)
    assert Student.objects.count() == 1


def test_timetable_overlap_with_different_period(erp):
    first = Period.objects.create(school=erp.school, name="One", order=1, start_time=time(9), end_time=time(10))
    overlap = Period.objects.create(
        school=erp.school, name="Two", order=2, start_time=time(9, 30), end_time=time(10, 30)
    )
    RoutineSlot.objects.create(
        school=erp.school,
        academic_year=erp.year,
        section=erp.section,
        weekday=1,
        period=first,
        subject=erp.subject,
        teacher=erp.employee,
    )
    slot = RoutineSlot(
        school=erp.school,
        academic_year=erp.year,
        section=erp.other_section,
        weekday=1,
        period=overlap,
        subject=erp.subject,
        teacher=erp.employee,
    )
    with pytest.raises(ValidationError):
        slot.full_clean()


def test_sms_queues_and_deduplicates_without_sending(erp):
    batch = queue_batch(
        erp.school, "Test", "custom", "Hello", [("One", "01712345678"), ("Same", "+8801712345678")], erp.admin
    )
    assert batch.total == 1
    assert SMSMessage.objects.get().status == "queued"


def test_sms_invalid_number_rolls_back(erp):
    with pytest.raises(ValidationError):
        queue_batch(erp.school, "Test", "custom", "Hello", [("Bad", "123")], erp.admin)
    assert not SMSMessage.objects.exists()


def test_sms_failure_logged(erp):
    batch = queue_batch(erp.school, "Test", "custom", "Hello", [("One", "01712345678")], erp.admin)

    class Failure(BaseSMSBackend):
        def send(self, message):
            raise RuntimeError("Gateway down")

    message = batch.messages.get()
    assert Failure().deliver(message) is False
    message.refresh_from_db()
    assert message.status == "failed"


def test_result_exports_pdf_and_excel(admin_client, erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    exam = publish_exam(erp.exam, erp.admin)
    for fmt in ("csv", "xlsx", "pdf"):
        response = admin_client.get(f"/exams/results/?exam={exam.pk}&class_level={erp.level.pk}&format={fmt}")
        assert response.status_code == 200
        assert response.content
    response = admin_client.get(f"/exams/{exam.pk}/report/{erp.student.pk}/?format=pdf")
    assert response.content.startswith(b"%PDF")
    from openpyxl import load_workbook

    response = admin_client.get(f"/exams/results/?exam={exam.pk}&class_level={erp.level.pk}&format=xlsx")
    assert load_workbook(BytesIO(response.content)).sheetnames == ["results", "Subject analysis", "About"]


def test_parent_cannot_read_other_student_or_draft_results(client, erp, invoice):
    client.force_login(erp.parent)
    other = Student.objects.create(
        school=erp.school,
        student_id="S2",
        first_name="Private",
        gender="M",
        date_of_birth=date(2016, 1, 1),
        admission_date=date(2026, 1, 1),
    )
    assert client.get(f"/students/{other.pk}/").status_code == 404
    assert client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/").status_code == 403
    assert client.post(f"/fees/{invoice.pk}/", {"amount": 1}).status_code == 403


def test_cross_school_object_edit_denied(admin_client, erp):
    other = Student.objects.create(
        school=erp.other,
        student_id="S1",
        first_name="Private",
        gender="M",
        date_of_birth=date(2016, 1, 1),
        admission_date=date(2026, 1, 1),
    )
    assert admin_client.get(f"/students/{other.pk}/edit/").status_code == 404


def test_user_password_hashed_and_roles(admin_client, erp):
    group = Group.objects.get(name="Student")
    r = admin_client.post(
        "/users/new/",
        {"username": "new-user", "password": "Unique-Strong-249!", "roles": [group.pk], "is_active": "on"},
    )
    assert r.status_code == 302
    from users.models import User

    user = User.objects.get(username="new-user")
    assert user.school == erp.school and user.check_password("Unique-Strong-249!")
    assert not user.is_superuser and not user.is_staff


def test_csrf_is_enforced(erp):
    from django.test import Client

    c = Client(enforce_csrf_checks=True)
    c.force_login(erp.admin)
    assert c.post("/users/new/", {}).status_code == 403


def test_published_mark_database_trigger_blocks_update_delete_insert(erp):
    from django.db import DatabaseError, transaction

    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(70))
    publish_exam(erp.exam, erp.admin)
    for operation in (
        lambda: Mark.objects.update(marks_obtained=99),
        lambda: Mark.objects.all().delete(),
        lambda: Mark.objects.create(
            school=erp.school, schedule=erp.schedule, enrollment=erp.enrollment, marks_obtained=99
        ),
    ):
        with pytest.raises(DatabaseError):
            with transaction.atomic():
                operation()
    assert Mark.objects.get().marks_obtained == 70


def test_mark_override_does_not_leak_after_correction(erp):
    from django.db import DatabaseError, transaction

    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(70))
    publish_exam(erp.exam, erp.admin)
    unlock = UnlockRequest.objects.create(
        school=erp.school, schedule=erp.schedule, requested_by=erp.teacher, reason="Correction"
    )
    review_unlock(unlock, erp.admin, True)
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80), expected_version=1)
    with pytest.raises(DatabaseError):
        with transaction.atomic():
            Mark.objects.update(marks_obtained=99)


def test_published_exam_schedule_edit_rejected(admin_client, erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(70))
    publish_exam(erp.exam, erp.admin)
    r = admin_client.post(
        f"/exams/schedule/{erp.schedule.pk}/edit/",
        {
            "exam": erp.exam.pk,
            "class_level": erp.level.pk,
            "subject": erp.subject.pk,
            "full_marks": 50,
            "pass_marks": 20,
        },
    )
    assert r.status_code == 200 and b"published exams cannot be changed" in r.content
    erp.schedule.refresh_from_db()
    assert erp.schedule.full_marks == 100


def test_multicategory_payment_allocation(erp, invoice):
    from fees.models import FeeCategory, FeeInvoiceItem

    tuition = Account.objects.get(school=erp.school, code="4010")
    exam_account = Account.objects.get(school=erp.school, code="4030")
    erp.category.income_account = tuition
    erp.category.save()
    cat = FeeCategory.objects.create(school=erp.school, name="Exam", income_account=exam_account)
    FeeInvoiceItem.objects.create(invoice=invoice, category=cat, amount=500)
    payment = pay(erp, invoice, 300)
    assert tuition.balance() == 200 and exam_account.balance() == 100
    assert payment.journal_entry.is_balanced
    cancel_payment(school=erp.school, user=erp.accountant, payment=payment, reason="Correction")
    assert tuition.balance() == 0 and exam_account.balance() == 0


def test_download_permissions_and_private_media(client, erp, settings, tmp_path):
    from downloads.models import DownloadCategory, DownloadItem

    settings.MEDIA_ROOT = tmp_path
    cat = DownloadCategory.objects.create(school=erp.school, name="Private")
    item = DownloadItem.objects.create(
        school=erp.school,
        category=cat,
        title="Staff only",
        audience="staff",
        file=SimpleUploadedFile("staff.txt", b"private staff content"),
    )
    client.force_login(erp.parent)
    assert client.get(f"/downloads/{item.pk}/file/").status_code == 403
    assert client.get("/media/" + item.file.name).status_code == 404
    client.force_login(erp.teacher)
    response = client.get(f"/downloads/{item.pk}/file/")
    assert response.status_code == 200
    assert b"".join(response.streaming_content) == b"private staff content"
    item.refresh_from_db()
    assert item.download_count == 1


def test_create_student_and_duplicate_validation(admin_client, erp):
    data = {
        "student_id": "S2",
        "first_name": "New",
        "gender": "M",
        "date_of_birth": "2016-01-01",
        "admission_date": "2026-01-01",
        "status": "active",
    }
    response = admin_client.post("/students/new/", data)
    assert response.status_code == 302
    assert Student.objects.get(student_id="S2").school == erp.school
    response = admin_client.post("/students/new/", data)
    assert response.status_code == 200 and b"already exists" in response.content
    assert Student.objects.filter(student_id="S2").count() == 1


def test_user_edit_cannot_escalate_to_platform_superuser(admin_client, erp):
    response = admin_client.post(
        f"/users/{erp.teacher.pk}/edit/",
        {
            "username": "teacher",
            "roles": [Group.objects.get(name="Teacher").pk],
            "is_active": "on",
            "is_superuser": "on",
            "is_staff": "on",
        },
    )
    assert response.status_code == 302
    erp.teacher.refresh_from_db()
    assert not erp.teacher.is_superuser and not erp.teacher.is_staff


def test_student_post_cannot_link_cross_school_user(admin_client, erp):
    from users.models import User

    user = User.objects.create_user("foreign", school=erp.other)
    response = admin_client.post(
        "/students/new/",
        {
            "student_id": "S2",
            "user": user.pk,
            "first_name": "New",
            "gender": "M",
            "date_of_birth": "2016-01-01",
            "admission_date": "2026-01-01",
            "status": "active",
        },
    )
    assert response.status_code == 200
    assert b"Select a valid choice" in response.content
    assert not Student.objects.filter(student_id="S2").exists()


def test_payroll_posts_once_and_receipt_pdf(admin_client, erp, invoice):
    from finance.models import Payroll

    row = Payroll.objects.create(
        school=erp.school, employee=erp.employee, month=date(2026, 9, 1), basic=1000, allowance=100, deduction=50
    )
    for _ in range(2):
        assert admin_client.post(f"/finance/payroll/{row.pk}/pay/").status_code == 302
    assert JournalEntry.objects.filter(source="salary").count() == 1
    assert Account.objects.get(school=erp.school, code="5010").balance() == 1050
    assert admin_client.get(f"/finance/payroll/{row.pk}.pdf").content.startswith(b"%PDF")
    payment = pay(erp, invoice)
    assert admin_client.get(f"/fees/receipts/{payment.pk}.pdf").content.startswith(b"%PDF")


def test_seed_is_idempotent_and_sends_no_messages(erp):
    from django.core.management import call_command

    call_command("seed_demo")
    count = Student.objects.count()
    call_command("seed_demo")
    assert Student.objects.count() == count
    assert SMSMessage.objects.count() == 0


def test_invalid_mark_id_returns_400(admin_client, erp):
    response = admin_client.post(
        "/exams/marks/save/",
        {
            "schedule": erp.schedule.pk,
            "section": erp.section.pk,
            "enrollment": "bad",
            "score": 50,
            "expected_version": 0,
        },
    )
    assert response.status_code == 400


def test_login_rate_limit(client, erp):
    from django.core.cache import cache

    cache.clear()
    for _ in range(5):
        client.post("/login/", {"username": "admin", "password": "wrong"})
    response = client.post("/login/", {"username": "admin", "password": "Test-pass-9842"})
    assert response.status_code == 200 and b"Too many attempts" in response.content
    cache.clear()
