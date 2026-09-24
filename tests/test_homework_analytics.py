"""
Phase 6e: homework analytics.

Rates count only work whose outcome is known: absent and excused are left out, and work checked
in class that nobody has checked yet is counted apart, never as missing. A family sees their own
child against the section as a whole; a subject teacher, their subjects and sections; a class
teacher, their section; the managers, the school. None of it exists for a school without the
module, and early warning mentions homework only for a school that has it.
"""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.test import Client
from django.utils import timezone

from homework import services
from homework.analytics import base_rows, summarise
from homework.models import Task
from students.models import Enrollment, Student


def login(user):
    client = Client()
    client.force_login(user)
    return client


def classmate(erp, name="Karim", roll=2):
    person = Student.objects.create(
        school=erp.school,
        student_id=f"S{roll}",
        first_name=name,
        gender="M",
        date_of_birth=date(2016, 1, 1),
        admission_date=date(2026, 1, 1),
    )
    return Enrollment.objects.create(
        school=erp.school,
        student=person,
        academic_year=erp.year,
        class_level=erp.level,
        section=erp.section,
        roll_number=roll,
    )


def task_with(erp, statuses, *, days_ago=1, hand_in=Task.HandIn.IN_CLASS, marks=None, title="Work"):
    """A task due `days_ago` days ago, with each enrollment's record set as given."""
    now = timezone.now()
    task = Task(
        school=erp.school,
        academic_year=erp.year,
        class_level=erp.level,
        subject=erp.subject,
        title=title,
        hand_in=hand_in,
        marking=Task.Marking.MARKS if marks else Task.Marking.NONE,
        max_marks=Decimal(10) if marks else None,
    )
    task = services.save_task(
        user=erp.teacher,
        task=task,
        targets={erp.section: now - timedelta(days=days_ago)},
        action="publish",
        now=now - timedelta(days=days_ago + 2),
    )
    rows = []
    for enrollment, status in statuses.items():
        if status == "pending":
            continue
        row = task.submissions.get(enrollment=enrollment)
        if status == "family":
            services.family_mark_done(
                user=erp.parent, submission=row, done=True, now=now - timedelta(days=days_ago + 1)
            )
            continue
        rows.append(
            {
                "id": row.pk,
                "version": row.version,
                "status": "done" if status == "late" else status,
                "late": status == "late",
                "mark": (marks or {}).get(enrollment.pk),
                "grade": "",
                "feedback": "",
                "reason": "Ill" if status == "excused" else "",
            }
        )
    if rows:
        services.check_rows(
            user=erp.teacher, task=task, target=task.targets.get(), rows=rows, return_work=True, now=now
        )
    return task


def test_the_rates_count_only_work_whose_outcome_is_known(erp, homework):
    others = [classmate(erp, f"Pupil {n}", n) for n in range(2, 8)]
    everyone = [erp.enrollment, *others]
    statuses = ["done", "partial", "not_done", "absent", "excused", "pending", "late"]
    task_with(erp, dict(zip(everyone, statuses, strict=True)))
    now = timezone.now()
    figures = summarise(base_rows(erp.school, erp.year, now - timedelta(weeks=4), now))
    # Seven records: absent and excused left out, the unchecked one counted apart.
    assert figures["known"] == 4 and figures["unchecked"] == 1
    assert figures["handed"] == 3 and figures["rate"] == 75  # done, partly done, late
    assert figures["on_time_rate"] == 50 and figures["missing"] == 1


def test_online_work_not_handed_in_is_missing_and_a_family_tick_is_its_own_share(erp, homework):
    second = classmate(erp)
    task_with(erp, {erp.enrollment: "pending", second: "pending"}, hand_in=Task.HandIn.ONLINE)
    task_with(erp, {erp.enrollment: "family", second: "done"}, title="Ticked")
    now = timezone.now()
    figures = summarise(base_rows(erp.school, erp.year, now - timedelta(weeks=4), now))
    assert figures["known"] == 4 and figures["handed"] == 2 and figures["missing"] == 2
    assert figures["family_share"] == 50


def test_marks_are_averaged_as_percentages_of_returned_work(erp, homework):
    second = classmate(erp)
    task_with(
        erp, {erp.enrollment: "done", second: "done"}, marks={erp.enrollment.pk: Decimal(8), second.pk: Decimal(5)}
    )
    now = timezone.now()
    assert summarise(base_rows(erp.school, erp.year, now - timedelta(weeks=4), now))["mark"] == Decimal("65.0")


# ------------------------------------------------------------------ who sees what


def test_a_family_sees_their_child_against_the_section_and_no_other_name(erp, homework):
    second = classmate(erp, "Karim")
    task_with(erp, {erp.enrollment: "not_done", second: "done"}, title="Tables")
    page = login(erp.parent).get("/portal/progress/").content.decode()
    # No exam is published, yet the homework shows.
    assert "Handed in" in page and "Tables" in page and "50%" in page
    assert "Karim" not in page


def test_staff_open_a_students_homework_without_any_published_exam(erp, homework):
    task_with(erp, {erp.enrollment: "done"})
    page = login(erp.teacher).get(f"/analytics/student/{erp.student.pk}/")
    assert page.status_code == 200 and "On time" in page.content.decode()


def test_each_teacher_sees_their_own_subjects_and_sections(erp, homework):
    from tests.test_homework_setting import _colleague

    task_with(erp, {erp.enrollment: "done"}, title="Section A work")
    mine = login(erp.teacher).get("/homework/insights/").content.decode()
    assert "The subjects you teach" in mine and "Math" in mine
    colleague = _colleague(erp)
    theirs = login(colleague).get("/homework/insights/").content.decode()
    assert "The subjects you teach" not in theirs
    url = f"/homework/insights/subject/?section={erp.section.pk}&subject={erp.subject.pk}"
    assert login(colleague).get(url).status_code == 403
    grid = login(erp.teacher).get(url).content.decode()
    assert "Ayesha" in grid and "Section A work" in grid


def test_a_class_teacher_sees_their_section_across_subjects(erp, homework):
    task_with(erp, {erp.enrollment: "done"})
    url = f"/homework/insights/section/?section={erp.section.pk}"
    assert login(erp.teacher).get(url).status_code == 403
    erp.section.class_teacher = erp.employee
    erp.section.save()
    page = login(erp.teacher).get(url).content.decode()
    assert "Ayesha" in page and "Math" in page


def test_the_managers_see_the_school_and_can_take_it_to_excel(erp, homework):
    from openpyxl import load_workbook

    task_with(erp, {erp.enrollment: "done"})
    page = login(erp.admin).get("/homework/insights/").content.decode()
    assert "By class" in page and "Teachers" in page and "Class 1" in page
    book = load_workbook(BytesIO(login(erp.admin).get("/homework/insights/?format=xlsx").content))
    assert book.sheetnames[:5] == ["homework-4w", "Sections", "Subjects", "Teachers", "Load"]


def test_early_warning_mentions_homework_only_with_the_module(erp, homework):
    for n in range(4):
        task_with(erp, {erp.enrollment: "not_done"}, title=f"Missed {n}", days_ago=n + 1)
    url = f"/reports/early-warning/?section={erp.section.pk}"
    body = login(erp.admin).get(url).content.decode()
    assert "Homework: 0 of 4 handed in" in body
    erp.school.homework_enabled = False
    erp.school.save()
    body = login(erp.admin).get(url).content.decode()
    assert "Homework" not in body


def test_nothing_of_it_shows_without_the_module(erp):
    assert "Homework analytics" not in login(erp.admin).get("/analytics/").content.decode()
    erp.school.homework_enabled = True
    erp.school.save()
    assert "Homework analytics" in login(erp.admin).get("/analytics/").content.decode()


def test_the_school_figures_take_a_steady_number_of_queries(erp, homework):
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    from homework.analytics import school_overview

    def queries():
        now = timezone.now()
        with CaptureQueriesContext(connection) as captured:
            school_overview(erp.school, erp.year, now - timedelta(weeks=4), now)
        return len(captured)

    task_with(erp, {erp.enrollment: "done"}, title="First")
    small = queries()
    others = [classmate(erp, f"Pupil {n}", n) for n in range(2, 12)]
    for n in range(3):
        task_with(erp, dict.fromkeys([erp.enrollment, *others], "done"), title=f"Task {n}", days_ago=n + 1)
    assert queries() == small
