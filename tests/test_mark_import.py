"""
Exam-hall sheets and marks from a spreadsheet.

The collection sheet and marks register print the paper's parts for the students who sit it.
An import is checked before anything is saved, matched by student ID or roll and never by
name, leaves blank rows alone, and saves through the mark grid's all-or-nothing path.
"""

import io
from datetime import date
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import HttpResponse
from django.test import Client
from openpyxl import Workbook, load_workbook

import examinations.documents as documents
from examinations.models import Mark, PaperComponent
from examinations.services import publish_exam, save_mark
from students.models import Enrollment, Student
from tests.test_rulebooks import card_texts


def login(user):
    client = Client()
    client.force_login(user)
    return client


@pytest.fixture
def pair(erp):
    """Ayesha (S1, roll 1) and Babul (S2, roll 2) in Class 1 A, which the teacher teaches Math."""
    student = Student.objects.create(
        school=erp.school,
        student_id="S2",
        first_name="Babul",
        gender="M",
        date_of_birth=date(2016, 1, 1),
        admission_date=date(2026, 1, 1),
    )
    second = Enrollment.objects.create(
        school=erp.school,
        student=student,
        academic_year=erp.year,
        class_level=erp.level,
        section=erp.section,
        roll_number=2,
    )
    return erp, second


def query(erp, **extra):
    return "&".join(f"{k}={v}" for k, v in {"schedule": erp.schedule.pk, "section": erp.section.pk, **extra}.items())


def sheet_text(client, erp, monkeypatch, **extra):
    captured = {}

    def fake_document(school, title, flowables, **kwargs):
        captured.update(title=title, text=card_texts(flowables), **kwargs)
        return HttpResponse(b"%PDF-fake", content_type="application/pdf")

    monkeypatch.setattr(documents, "document", fake_document)
    response = client.get(f"/exams/marks/sheet.pdf?{query(erp, **extra)}")
    assert response.status_code == 200
    return captured


# ------------------------------------------------------------------ exam-hall sheets


def test_the_blank_sheet_has_a_column_for_each_part_and_a_row_for_each_student(pair, monkeypatch):
    erp, _second = pair
    PaperComponent.objects.create(
        school=erp.school, schedule=erp.schedule, code="cq", name="Creative", full_marks=70, pass_marks=23, order=1
    )
    PaperComponent.objects.create(
        school=erp.school,
        schedule=erp.schedule,
        code="mcq",
        name="Multiple choice",
        full_marks=30,
        pass_marks=10,
        order=2,
    )
    sheet = sheet_text(login(erp.teacher), erp, monkeypatch)
    assert sheet["title"] == "Mark collection sheet"
    for heading in ("Creative (70)", "Multiple choice (30)", "Total (100)", "Remarks", "Examiner"):
        assert heading in sheet["text"], heading
    assert "Ayesha" in sheet["text"] and "Babul" in sheet["text"]


def test_the_marks_register_prints_what_was_entered(pair, monkeypatch):
    erp, second = pair
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal("84.50"))
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=second, absent=True)
    sheet = sheet_text(login(erp.teacher), erp, monkeypatch, kind="register")
    assert sheet["title"] == "Marks register"
    assert "84.5" in sheet["text"] and "ABS" in sheet["text"]


def test_sheets_are_only_for_the_teacher_of_that_section(pair):
    erp, _second = pair
    other = f"/exams/marks/sheet.pdf?schedule={erp.schedule.pk}&section={erp.other_section.pk}"
    assert login(erp.teacher).get(other).status_code in (403, 404)
    assert login(erp.parent).get(f"/exams/marks/sheet.pdf?{query(erp)}").status_code == 403


# ------------------------------------------------------------------ import


def workbook_upload(rows, name="marks.xlsx"):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def check(client, erp, upload):
    return client.post(
        f"/exams/marks/import/?{query(erp)}",
        {**dict(p.split("=") for p in query(erp).split("&")), "action": "check", "file": upload},
    )


def confirm(client, erp):
    return client.post(
        "/exams/marks/import/", {**dict(p.split("=") for p in query(erp).split("&")), "action": "confirm"}
    )


def test_the_template_round_trips_into_saved_marks(pair):
    erp, second = pair
    client = login(erp.teacher)
    download = client.get(f"/exams/marks/import/?{query(erp, download='xlsx')}")
    sheet = load_workbook(io.BytesIO(download.content)).worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert rows[0] == ["Roll", "Student ID", "Student", "Marks (100)"]
    assert [r[1] for r in rows[1:]] == ["S1", "S2"]
    rows[1][3], rows[2][3] = 72, "abs"

    page = check(client, erp, workbook_upload(rows)).content.decode()
    assert "2 new" in page and "Save 2 marks" in page
    assert not Mark.objects.exists()  # nothing saved by the check

    assert confirm(client, erp).status_code == 302
    assert Mark.objects.get(enrollment=erp.enrollment).marks_obtained == Decimal("72.00")
    assert Mark.objects.get(enrollment=second).is_absent


def test_a_blank_row_leaves_the_mark_alone_and_rows_match_by_roll_too(pair):
    erp, second = pair
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=second, score=Decimal(55))
    client = login(erp.teacher)
    rows = [["Roll", "Student ID", "Student", "Marks (100)"], [1, "", "whatever name", 64], [2, "S2", "Babul", ""]]
    page = check(client, erp, workbook_upload(rows)).content.decode()
    assert "1 new" in page and "0 changed" in page
    confirm(client, erp)
    assert Mark.objects.get(enrollment=erp.enrollment).marks_obtained == Decimal(64)
    assert Mark.objects.get(enrollment=second).marks_obtained == Decimal(55)


def test_problems_are_listed_and_nothing_can_be_saved(pair):
    erp, _second = pair
    client = login(erp.teacher)
    rows = [
        ["Roll", "Student ID", "Student", "Marks (100)"],
        [1, "S1", "Ayesha", 120],
        [9, "NOPE", "Someone", 50],
        [2, "S2", "Babul", "seventy"],
    ]
    page = check(client, erp, workbook_upload(rows)).content.decode()
    assert "more than the full marks, 100" in page
    assert "no student with student ID NOPE" in page
    assert "&#x27;seventy&#x27; is not a number" in page
    assert "Save " not in page
    confirm(client, erp)
    assert not Mark.objects.exists()


def test_parts_are_imported_and_must_all_be_filled(pair):
    erp, second = pair
    PaperComponent.objects.create(
        school=erp.school, schedule=erp.schedule, code="cq", name="Creative", full_marks=70, pass_marks=23, order=1
    )
    PaperComponent.objects.create(
        school=erp.school,
        schedule=erp.schedule,
        code="mcq",
        name="Multiple choice",
        full_marks=30,
        pass_marks=10,
        order=2,
    )
    client = login(erp.teacher)
    rows = [
        ["Roll", "Student ID", "Student", "Creative (70)", "Multiple choice (30)"],
        [1, "S1", "Ayesha", 50, 22],
        [2, "S2", "Babul", 40, ""],
    ]
    page = check(client, erp, workbook_upload(rows)).content.decode()
    assert "fill in every part" in page
    rows[2][4] = 18
    check(client, erp, workbook_upload(rows))
    confirm(client, erp)
    mark = Mark.objects.get(enrollment=erp.enrollment)
    assert mark.marks_obtained == Decimal(72) and mark.component_marks == {"cq": "50.00", "mcq": "22.00"}
    assert Mark.objects.get(enrollment=second).marks_obtained == Decimal(58)


def test_a_csv_template_with_its_notes_above_the_table_reads_back(pair):
    erp, _second = pair
    client = login(erp.teacher)
    text = client.get(f"/exams/marks/import/?{query(erp, download='csv')}").content.decode("utf-8-sig")
    text = text.replace("S1,Ayesha,", "S1,Ayesha,81")
    page = check(client, erp, SimpleUploadedFile("marks.csv", text.encode("utf-8"))).content.decode()
    assert "1 new" in page


def test_a_mark_changed_by_someone_else_after_the_check_is_not_overwritten(pair):
    erp, _second = pair
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(40))
    client = login(erp.teacher)
    check(client, erp, workbook_upload([["Roll", "Student ID", "Student", "Marks (100)"], [1, "S1", "Ayesha", 45]]))
    save_mark(user=erp.admin, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(48), expected_version=1)
    response = confirm(client, erp)
    assert response.status_code == 200 and "Nothing was saved" in response.content.decode()
    assert Mark.objects.get(enrollment=erp.enrollment).marks_obtained == Decimal(48)


def test_published_marks_cannot_be_imported_and_other_teachers_cannot_import(pair):
    erp, second = pair
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(70))
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=second, score=Decimal(60))
    other = f"/exams/marks/import/?schedule={erp.schedule.pk}&section={erp.other_section.pk}"
    assert login(erp.teacher).get(other).status_code in (403, 404)
    publish_exam(erp.exam, erp.admin)
    assert login(erp.teacher).get(f"/exams/marks/import/?{query(erp)}").status_code == 403
