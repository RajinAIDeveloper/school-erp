"""
Rulebooks: each programme's results worked out by its own rules, and no other's.

Cambridge and Edexcel give a grade per subject and nothing more; the IB MYP turns four
criteria into a grade from 1 to 7; the Bangladesh board rules apply only where chosen; and a
school that never chooses keeps its own rules. The thresholds in these examples are typical
internal ones, as a school would set them, not an awarding body's official thresholds.
"""

from datetime import date
from decimal import Decimal

import pytest
from django.core.exceptions import ValidationError
from django.test import Client

from academics.models import ClassLevel, Section, Subject, SubjectTeacher
from core.models import AssessmentSystem
from examinations.grading import combine_units, grade_paper, grades_outcome, headline, myp_outcome
from examinations.models import Exam, ExamSchedule, Mark, PaperComponent, ResultSnapshot
from examinations.presets import PRESETS, install_preset
from examinations.rulebooks import RULEBOOKS, rulebook
from examinations.services import build_result_sheet, publish_exam, save_mark
from students.models import Enrollment, Student

# ================================================================ presets


@pytest.mark.parametrize("key", list(PRESETS))
def test_every_preset_installs_a_complete_ordered_scale(erp, key):
    scale, created = install_preset(erp.school, key)
    assert created
    rules = list(scale.rules.order_by("-min_percent"))
    assert rules[0].max_percent == 100 and rules[-1].min_percent == 0
    for higher, lower in zip(rules, rules[1:], strict=False):
        # No gaps and no overlaps between neighbouring bands.
        assert lower.max_percent == higher.min_percent - Decimal("0.01")


def test_installing_a_preset_twice_leaves_the_school_s_edits_alone(erp):
    scale, _ = install_preset(erp.school, "cambridge-igcse")
    top = scale.rules.get(letter="A*")
    top.min_percent = Decimal("92")
    top.save()
    again, created = install_preset(erp.school, "cambridge-igcse")
    assert not created and again.pk == scale.pk
    assert again.rules.get(letter="A*").min_percent == Decimal("92")


def test_the_presets_screen_adds_a_scale(erp):
    client = Client()
    client.force_login(erp.admin)
    assert client.get("/exams/scales/presets/").status_code == 200
    response = client.post("/exams/scales/presets/", {"preset": "igcse-9-1"})
    assert response.status_code == 302
    assert erp.school.pk in {
        s.school_id for s in erp.scale.__class__.objects.filter(name__startswith="International GCSE")
    }


def test_a_teacher_cannot_add_scales(erp):
    client = Client()
    client.force_login(erp.teacher)
    assert client.post("/exams/scales/presets/", {"preset": "igcse-9-1"}).status_code == 403


# ================================================================ the arithmetic


def rules_of(erp, key):
    from examinations.grading import scale_rules

    scale, _ = install_preset(erp.school, key)
    return scale_rules(scale)


def spec(name, full=100, pass_marks=33):
    return {
        "schedule_id": hash(name) % 10000,
        "subject": name,
        "subject_id": 1,
        "subject_code": "",
        "unit_id": hash(name) % 10000,
        "unit_name": name,
        "full_marks": Decimal(full),
        "pass_marks": Decimal(pass_marks),
        "components": [],
        "role": "main",
    }


def test_cambridge_has_no_pass_mark_a_low_score_is_just_a_low_grade(erp):
    """25% under a 33% pass mark would be a fail on the national scale; on IGCSE it is a G."""
    rules = rules_of(erp, "cambridge-igcse")
    cell = grade_paper(spec("Physics"), {"absent": False, "score": Decimal(25), "parts": {}}, rules, pass_marks=False)
    assert cell["letter"] == "G"
    cell = grade_paper(spec("Physics"), {"absent": False, "score": Decimal(12), "parts": {}}, rules, pass_marks=False)
    assert cell["letter"] == "U"


def test_an_absence_under_cambridge_has_no_grade(erp):
    rules = rules_of(erp, "cambridge-igcse")
    cell = grade_paper(spec("Physics"), {"absent": True, "score": None, "parts": {}}, rules, pass_marks=False)
    assert cell["letter"] == "ABS"


def test_grades_outcome_reports_a_tally_and_no_gpa_or_result(erp):
    rules = rules_of(erp, "cambridge-igcse")
    cells = [
        grade_paper(spec(name), {"absent": False, "score": Decimal(score), "parts": {}}, rules, pass_marks=False)
        for name, score in (("Physics", 95), ("Chemistry", 91), ("Biology", 84), ("English", 72))
    ]
    outcome = grades_outcome(combine_units(cells, rules, combine=False), rules)
    assert outcome["gpa"] is None and outcome["result"] is None
    assert outcome["headline"] == "2 A*, 1 A, 1 B"


def test_edexcel_9_to_1(erp):
    rules = rules_of(erp, "igcse-9-1")
    cell = grade_paper(spec("Maths"), {"absent": False, "score": Decimal(47), "parts": {}}, rules, pass_marks=False)
    assert cell["letter"] == "4"  # a standard pass on the 9-1 scale


@pytest.mark.parametrize(
    "criteria,grade",
    [
        ((6, 7, 5, 6), "6"),
        ((6, 6, 5, 6), "5"),
        ((8, 8, 6, 6), "7"),
        ((7, 7, 7, 6), "6"),
        ((1, 2, 1, 1), "1"),
        ((2, 2, 1, 1), "2"),
    ],
)
def test_myp_criteria_total_converts_by_the_ib_boundaries(erp, criteria, grade):
    """24 of 32 is a 6 and 23 a 5; 28 a 7; 5 a 1 and 6 a 2."""
    rules = rules_of(erp, "ib-myp")
    total = sum(criteria)
    cell = grade_paper(
        spec("Sciences", full=32), {"absent": False, "score": Decimal(total), "parts": {}}, rules, pass_marks=False
    )
    assert cell["letter"] == grade


def test_myp_outcome_sums_the_subject_grades(erp):
    rules = rules_of(erp, "ib-myp")
    cells = [
        grade_paper(
            spec(name, full=32), {"absent": False, "score": Decimal(score), "parts": {}}, rules, pass_marks=False
        )
        for name, score in (("Sciences", 24), ("Mathematics", 28), ("Language and literature", 19))
    ]
    outcome = myp_outcome(combine_units(cells, rules, combine=False), rules)
    assert outcome["points"] == 6 + 7 + 5
    assert outcome["gpa"] is None and outcome["result"] is None


def test_an_unknown_rulebook_is_refused_not_guessed():
    with pytest.raises(ValidationError):
        rulebook("made-up")


def test_only_the_national_rulebook_has_the_fourth_subject_and_paper_combining():
    for key, book in RULEBOOKS.items():
        assert book.fourth_subject == (key == "national")
        assert book.combine_papers == (key == "national")


# ================================================================ end to end: a Cambridge class


@pytest.fixture
def igcse(erp):
    """A Year 10 IGCSE class under Cambridge rules, with a weighted Physics paper."""
    scale, _ = install_preset(erp.school, "cambridge-igcse")
    level = ClassLevel.objects.create(
        school=erp.school, name="Year 10", order=10, assessment_system=AssessmentSystem.CAMBRIDGE
    )
    section = Section.objects.create(school=erp.school, class_level=level, name="Blue")
    exam = Exam.objects.create(school=erp.school, academic_year=erp.year, name="Mock", grade_scale=scale)
    physics = Subject.objects.create(school=erp.school, name="Physics", code="0625")
    english = Subject.objects.create(school=erp.school, name="English as a Second Language", code="0510")
    papers = {}
    for subject in (physics, english):
        papers[subject.code] = ExamSchedule.objects.create(
            school=erp.school, exam=exam, class_level=level, subject=subject, full_marks=100, pass_marks=33
        )
        SubjectTeacher.objects.create(
            school=erp.school, academic_year=erp.year, section=section, subject=subject, teacher=erp.employee
        )
    # A weighted paper, as Cambridge IGCSE Physics 0625 is (June 2025 threshold table): Paper 2
    # multiple choice 40 raw at 30%, Paper 4 theory 80 raw at 50%, Paper 6 alternative to
    # practical 40 raw at 20%. Parts are scaled to their share, not added.
    for order, (code, name, full, weight) in enumerate(
        [("mcq", "Multiple choice", 40, 30), ("theory", "Theory", 80, 50), ("practical", "Practical", 40, 20)]
    ):
        PaperComponent.objects.create(
            school=erp.school,
            schedule=papers["0625"],
            code=code,
            name=name,
            full_marks=full,
            pass_marks=0,
            weight=weight,
            order=order,
        )
    pupils = []
    for roll, name in ((1, "Zara"), (2, "Imran")):
        student = Student.objects.create(
            school=erp.school,
            student_id=f"Y10-{roll}",
            first_name=name,
            gender="F",
            date_of_birth=date(2010, 5, 1),
            admission_date=date(2026, 1, 1),
        )
        pupils.append(
            Enrollment.objects.create(
                school=erp.school,
                student=student,
                academic_year=erp.year,
                class_level=level,
                section=section,
                roll_number=roll,
            )
        )
    from types import SimpleNamespace

    return SimpleNamespace(level=level, section=section, exam=exam, papers=papers, pupils=pupils, scale=scale)


def test_a_weighted_paper_is_scaled_not_added(erp, igcse):
    """
    32/40, 60/80 and 36/40 add to 128 of 160, which is 80% (an A). Scaled by weight it is
    0.80 x 30 + 0.75 x 50 + 0.90 x 20 = 79.5%, a B. The weighting changes the grade.
    """
    save_mark(
        user=erp.admin,
        schedule=igcse.papers["0625"],
        enrollment=igcse.pupils[0],
        components={"mcq": "32", "theory": "60", "practical": "36"},
    )
    assert Mark.objects.get(enrollment=igcse.pupils[0]).marks_obtained == Decimal("79.50")
    save_mark(user=erp.admin, schedule=igcse.papers["0510"], enrollment=igcse.pupils[0], score=Decimal(88))
    row = next(
        r for r in build_result_sheet(igcse.exam, igcse.level)["rows"] if r["enrollment_id"] == igcse.pupils[0].pk
    )
    physics = next(c for c in row["cells"] if c["subject"] == "Physics")
    assert physics["letter"] == "B"
    assert row["headline"] == "1 A, 1 B"
    assert row["gpa"] is None and row["result"] is None


def test_weights_must_cover_every_part_and_add_to_a_hundred(erp, igcse):
    PaperComponent.objects.filter(schedule=igcse.papers["0625"], code="practical").update(weight=10)
    with pytest.raises(ValidationError, match="100"):
        save_mark(
            user=erp.admin,
            schedule=igcse.papers["0625"],
            enrollment=igcse.pupils[0],
            components={"mcq": "32", "theory": "60", "practical": "36"},
        )


def mark_all(erp, igcse, scores):
    for pupil, (physics_parts, english) in zip(igcse.pupils, scores, strict=True):
        save_mark(user=erp.admin, schedule=igcse.papers["0625"], enrollment=pupil, components=physics_parts)
        save_mark(user=erp.admin, schedule=igcse.papers["0510"], enrollment=pupil, score=Decimal(english))


def test_a_cambridge_card_shows_grades_and_no_gpa_position_or_pass_fail(erp, igcse):
    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    publish_exam(igcse.exam, erp.admin)
    client = Client()
    client.force_login(erp.admin)
    body = client.get(f"/exams/{igcse.exam.pk}/report/{igcse.pupils[0].student_id}/").content.decode()
    assert "2 A*" in body
    assert ">GPA<" not in body
    assert "Rank in section" not in body
    assert "PASS" not in body and "FAIL" not in body
    assert "Worked out by" in body and "Cambridge" in body


def test_positions_can_be_turned_on_for_a_cambridge_exam(erp, igcse):
    igcse.exam.show_rank = True
    igcse.exam.save()
    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    publish_exam(igcse.exam, erp.admin)
    top = ResultSnapshot.objects.get(enrollment=igcse.pupils[0]).payload
    assert top["show_rank"] and top["rank"] == 1


def test_the_class_sheet_has_no_gpa_column_for_cambridge(erp, igcse):
    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    client = Client()
    client.force_login(erp.admin)
    query = f"exam={igcse.exam.pk}&class_level={igcse.level.pk}"
    body = client.get(f"/exams/results/?{query}").content.decode()
    assert "<th>GPA</th>" not in body and "Class rank" not in body
    assert "Overall" in body
    csv = client.get(f"/exams/results/?{query}&format=csv").content.decode("utf-8-sig")
    assert "GPA" not in csv.splitlines()[0] and "Class rank" not in csv.splitlines()[0]


def test_a_paper_can_use_another_board_s_scale(erp, igcse):
    """Edexcel 9-1 English in a Cambridge year: that paper is graded 9-1, the rest A*-G."""
    nine_to_one, _ = install_preset(erp.school, "igcse-9-1")
    igcse.papers["0510"].grade_scale = nine_to_one
    igcse.papers["0510"].save()
    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    row = next(
        r for r in build_result_sheet(igcse.exam, igcse.level)["rows"] if r["enrollment_id"] == igcse.pupils[0].pk
    )
    letters = {c["subject"]: c["letter"] for c in row["cells"]}
    assert letters["English as a Second Language"] == "9"
    assert letters["Physics"] == "A*"


def test_the_national_rules_never_reach_a_cambridge_class(erp, igcse):
    """Even with a 4th subject recorded, a Cambridge student gets no 4th-subject treatment."""
    igcse.pupils[0].fourth_subject = Subject.objects.get(code="0510")
    igcse.pupils[0].save()
    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    row = next(
        r for r in build_result_sheet(igcse.exam, igcse.level)["rows"] if r["enrollment_id"] == igcse.pupils[0].pk
    )
    assert not any(c["is_fourth"] for c in row["cells"])
    assert row["fourth_subject"] == ""


def test_the_trace_is_shown_to_staff_and_never_to_families(erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    publish_exam(erp.exam, erp.admin)
    client = Client()
    client.force_login(erp.admin)
    assert b"How this result was worked out" in client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/").content
    client.force_login(erp.parent)
    assert b"How this result was worked out" not in client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/").content


# ================================================================ backwards compatibility


def test_a_result_published_before_rulebooks_still_renders(erp):
    """An old snapshot has no headline, parts, attendance or rulebook fields; it must still print."""
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    publish_exam(erp.exam, erp.admin)
    snapshot = ResultSnapshot.objects.get()
    legacy = {
        key: snapshot.payload[key]
        for key in (
            "enrollment_id",
            "student_id",
            "student_code",
            "student",
            "section_id",
            "section",
            "class_level_id",
            "roll",
            "total",
            "full_total",
            "percent",
            "gpa",
            "result",
            "complete",
            "rank",
            "grade_rank",
        )
    }
    legacy["cells"] = [
        {k: v for k, v in cell.items() if k not in ("components", "is_fourth", "subject_code", "unit_id", "unit_name")}
        for cell in snapshot.payload["cells"]
    ]
    ResultSnapshot.objects.filter(pk=snapshot.pk).update(payload=legacy)
    assert headline(legacy) == f"GPA {legacy['gpa']} · PASS"
    client = Client()
    client.force_login(erp.admin)
    card = client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/")
    assert card.status_code == 200
    assert b"Rank in section" in card.content  # older results always showed positions
    assert client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/?format=pdf").content.startswith(b"%PDF")
    client.force_login(erp.parent)
    assert client.get("/portal/results/").status_code == 200


def test_a_school_that_never_chooses_gets_exactly_the_old_results(erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    row = build_result_sheet(erp.exam, erp.level)["rows"][0]
    assert row["system"] == "own"
    assert row["gpa"] == "5.00" and row["result"] == "PASS" and row["show_rank"]


def test_a_rulebook_without_pass_marks_ignores_part_pass_marks(erp):
    """
    A paper's parts carry their own pass marks. They must not switch pass-mark rules on for a
    rulebook that has none: an earlier version let the last part's pass mark do exactly that.
    """
    rules = rules_of(erp, "cambridge-igcse")
    paper = spec("Physics") | {"components": [("mcq", "MCQ", 40, 13), ("theory", "Theory", 60, 20)]}
    cell = grade_paper(
        paper, {"absent": False, "score": Decimal(30), "parts": {"mcq": "5", "theory": "25"}}, rules, pass_marks=False
    )
    assert cell["letter"] == "F"  # 30% is an F on the IGCSE scale, whatever the parts' pass marks
    assert cell["failed_part"] is False


def test_a_cambridge_card_says_it_is_the_schools_assessment_not_an_official_result(erp, igcse):
    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    publish_exam(igcse.exam, erp.admin)
    client = Client()
    client.force_login(erp.admin)
    url = f"/exams/{igcse.exam.pk}/report/{igcse.pupils[0].student_id}/"
    body = client.get(url).content.decode()
    assert "Cambridge grades (school assessment)" in body
    assert "Official results are issued only by Cambridge International Education." in body
    assert client.get(url + "?format=pdf").status_code == 200
    snapshot = ResultSnapshot.objects.filter(exam=igcse.exam).first()
    verify = Client().get(f"/exams/verify/{snapshot.verification_code}/").content.decode()
    assert "does not confirm any" in verify


def test_the_schools_own_rules_carry_no_awarding_body_notice(erp):
    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    client = Client()
    client.force_login(erp.admin)
    body = client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/").content.decode()
    assert "not an official result" not in body


# ================================================================ what a card prints


def card_texts(flowables):
    """Every piece of text in a card's PDF flowables, tables included."""
    from reportlab.platypus import Paragraph, Table

    out = []
    for item in flowables:
        if isinstance(item, Paragraph):
            out.append(item.getPlainText())
        elif isinstance(item, Table):
            out.extend(card_texts([cell for row in item._cellvalues for cell in row]))
        elif isinstance(item, list | tuple):
            out.extend(card_texts(item))
    return out


def test_a_cambridge_card_prints_no_points_column_and_marks_without_decimals(erp, igcse):
    from core.pdf import styles
    from examinations.documents import report_card_flowables

    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    publish_exam(igcse.exam, erp.admin)
    client = Client()
    client.force_login(erp.admin)
    body = client.get(f"/exams/{igcse.exam.pk}/report/{igcse.pupils[0].student_id}/").content.decode()
    assert "Points</th>" not in body
    assert '<td class="text-right">95</td>' in body and '<td class="text-right">100</td>' in body
    assert "95.00</td>" not in body and "100.00</td>" not in body
    assert "Multiple choice 38" in body
    assert "187 / 200" in body  # 95 + 92 of 200
    snapshot = ResultSnapshot.objects.get(exam=igcse.exam, enrollment=igcse.pupils[0])
    printed = card_texts(
        report_card_flowables(erp.school, igcse.exam, igcse.pupils[0], snapshot.payload, snapshot, styles())
    )
    assert "Points" not in printed
    assert "Total187 / 200" in printed and "95" in printed


def test_a_cambridge_results_page_has_no_failed_or_pass_rate(erp, igcse):
    import io

    from openpyxl import load_workbook

    mark_all(
        erp,
        igcse,
        [
            ({"mcq": "38", "theory": "76", "practical": "38"}, 92),
            ({"mcq": "20", "theory": "40", "practical": "20"}, 55),
        ],
    )
    client = Client()
    client.force_login(erp.admin)
    query = f"exam={igcse.exam.pk}&class_level={igcse.level.pk}"
    body = client.get(f"/exams/results/?{query}").content.decode()
    assert "Subject analysis" in body
    assert "Failed</th>" not in body and "Pass %" not in body
    workbook = load_workbook(io.BytesIO(client.get(f"/exams/results/?{query}&format=xlsx").content))
    headings = [cell.value for cell in workbook["Subject analysis"][1]]
    assert "Highest" in headings and "Failed" not in headings and "Pass percent" not in headings


def test_a_school_with_grade_points_keeps_the_points_and_pass_rate(erp):
    import io

    from openpyxl import load_workbook

    save_mark(user=erp.teacher, schedule=erp.schedule, enrollment=erp.enrollment, score=Decimal(80))
    client = Client()
    client.force_login(erp.admin)
    card = client.get(f"/exams/{erp.exam.pk}/report/{erp.student.pk}/").content.decode()
    assert "Points</th>" in card
    assert '<td class="text-right">80</td>' in card and "80.00</td>" not in card
    query = f"exam={erp.exam.pk}&class_level={erp.enrollment.class_level_id}"
    body = client.get(f"/exams/results/?{query}").content.decode()
    assert "Failed</th>" in body and "Pass %" in body
    workbook = load_workbook(io.BytesIO(client.get(f"/exams/results/?{query}&format=xlsx").content))
    assert "Pass percent" in [cell.value for cell in workbook["Subject analysis"][1]]


def test_dates_kept_as_text_print_like_every_other_date():
    from core.templatetags.erp import iso_date, mark

    assert iso_date("2026-09-01") == "01 Sep 2026"
    assert iso_date("") == ""
    assert mark("92.00") == "92" and mark("37.50") == "37.5" and mark("ABS") == "ABS" and mark(None) is None
