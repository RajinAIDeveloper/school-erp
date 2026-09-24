"""
Marks from a spreadsheet: a teacher downloads the class template, types the marks in Excel
(or anything that saves CSV), and uploads it.

Nothing is saved until the teacher has seen what the file would change. Rows are matched by
student ID, or by roll when the ID is left out, and never by name. A blank row leaves that
student's mark as it is, so a file with half the class changes only that half. The save goes
through the same all-or-nothing path as the mark grid: one bad row saves nothing, a mark
changed by someone else in the meantime is refused, and a published exam gets exactly one new
version.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError

from .parts import plain
from .services import save_marks

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
ABSENT_WORDS = {"abs", "a", "ab", "absent"}
EXEMPT_WORDS = {"ex", "exempt"}


# ------------------------------------------------------------------ the template


def headers_for(schedule, parts):
    heads = ["Roll", "Student ID", "Student"]
    if parts:
        return heads + [f"{part.name} ({plain(part.full_marks)})" for part in parts]
    return heads + [f"Marks ({plain(schedule.full_marks)})"]


def template_rows(parts, students, existing):
    """One row per student who sits the paper, with the marks already entered, so a file can go round trip."""
    rows = []
    for enrollment in students:
        line = [enrollment.roll_number, enrollment.student.student_id, enrollment.student.full_name]
        mark = existing.get(enrollment.pk)
        width = len(parts) or 1
        if mark is None:
            line += [""] * width
        elif mark.is_exempt:
            line += ["EX"] + [""] * (width - 1)
        elif mark.is_absent:
            line += ["ABS"] + [""] * (width - 1)
        elif parts:
            stored = mark.component_marks or {}
            line += [plain(stored[part.code]) if stored.get(part.code) not in (None, "") else "" for part in parts]
        else:
            line.append(plain(mark.marks_obtained))
        rows.append(line)
    return rows


# ------------------------------------------------------------------ reading a file


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_rows(upload):
    """The file's rows as lists of text, from an Excel workbook (.xlsx) or a CSV file."""
    if upload.size > MAX_BYTES:
        raise ValidationError("The file is larger than 2 MB. Upload the class template with the marks filled in.")
    data = upload.read()
    if upload.name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable workbook gets the same answer
            raise ValidationError("That Excel file could not be read. Save it as .xlsx and try again.") from exc
        rows = [[_text(cell) for cell in row] for row in workbook.worksheets[0].iter_rows(values_only=True)]
    elif upload.name.lower().endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        rows = [[_text(cell) for cell in row] for row in csv.reader(io.StringIO(text))]
    else:
        raise ValidationError("Upload the template as an Excel (.xlsx) or CSV file.")
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"The file has more than {MAX_ROWS} rows.")
    return rows


def _key(heading):
    """ "Creative (70)" and "creative" name the same column."""
    heading = heading.split("(")[0] if "(" in heading else heading
    return " ".join(heading.lower().replace("'", "").split())


# ------------------------------------------------------------------ checking


def check(schedule, parts, students, existing, rows):
    """
    What the file would change, and every problem with it. Writes nothing.

    Returns (prepared, problems, skipped): prepared is a list of dicts, one per student the
    file gives marks for, each with an action of new, changed or unchanged; problems are
    sentences naming the row; skipped explains rows deliberately left alone.
    """
    header_at = next(
        (i for i, row in enumerate(rows) if {"roll", "student id"} & {_key(cell) for cell in row if cell}), None
    )
    if header_at is None:
        raise ValidationError("No heading row with Roll or Student ID was found. Start from the class template.")
    heads = [_key(cell) for cell in rows[header_at]]

    def column(*names):
        return next((heads.index(name) for name in names if name in heads), None)

    id_col, roll_col = column("student id", "id"), column("roll")
    if parts:
        part_cols = {}
        for part in parts:
            at = column(_key(part.name), _key(part.code))
            if at is None:
                raise ValidationError(f"The file has no column for {part.name}. Start from the class template.")
            part_cols[part.code] = at
    else:
        marks_col = column("marks", "mark", "score")
        if marks_col is None:
            raise ValidationError("The file has no Marks column. Start from the class template.")

    by_id = {e.student.student_id.casefold(): e for e in students}
    by_roll = {str(e.roll_number): e for e in students}
    prepared, problems, skipped, seen = [], [], [], set()
    for number, row in enumerate(rows[header_at + 1 :], start=header_at + 2):

        def cell(at, row=row):
            return row[at] if at is not None and at < len(row) else ""

        student_id, roll = cell(id_col), cell(roll_col)
        values = [cell(at) for at in part_cols.values()] if parts else [cell(marks_col)]
        if not any(values):
            continue  # a blank row leaves that student's mark as it is
        enrollment = by_id.get(student_id.casefold()) if student_id else by_roll.get(roll)
        if enrollment is None:
            who = f"student ID {student_id}" if student_id else f"roll {roll or '(blank)'}"
            problems.append(f"Row {number}: no student with {who} sits this paper in this section.")
            continue
        if enrollment.pk in seen:
            problems.append(f"Row {number}: {enrollment.student.full_name} appears more than once.")
            continue
        seen.add(enrollment.pk)
        stored = existing.get(enrollment.pk)
        words = {value.lower() for value in values if value}
        if words & EXEMPT_WORDS or (stored and stored.is_exempt):
            skipped.append(f"Row {number}: {enrollment.student.full_name} is exempt; change that on the mark grid.")
            continue
        entry = {
            "row": number,
            "enrollment": enrollment.pk,
            "roll": enrollment.roll_number,
            "student": enrollment.student.full_name,
            "absent": bool(words & ABSENT_WORDS),
            "score": None,
            "parts": {},
            "version": stored.version if stored else 0,
        }
        error = None
        if not entry["absent"]:
            if parts:
                if not all(values):
                    error = "fill in every part, or leave them all blank"
                for part, value in zip(parts, values, strict=True):
                    if error:
                        break
                    error, amount = _number(value, part.full_marks, part.name)
                    entry["parts"][part.code] = str(amount) if amount is not None else ""
            else:
                error, amount = _number(values[0], schedule.full_marks, "Marks")
                entry["score"] = str(amount) if amount is not None else None
        if error:
            problems.append(f"Row {number} ({enrollment.student.full_name}): {error}.")
            continue
        entry["before"] = _shown(stored, parts)
        entry["after"] = _shown_entry(entry, parts)
        entry["action"] = "new" if stored is None else ("unchanged" if entry["before"] == entry["after"] else "changed")
        prepared.append(entry)
    return prepared, problems, skipped


def _number(value, full, label):
    try:
        amount = Decimal(value.replace(",", ""))
    except InvalidOperation:
        return f"{label}: '{value}' is not a number (write ABS for an absence)", None
    if not amount.is_finite() or amount < 0:
        return f"{label}: '{value}' is not a mark", None
    if amount > Decimal(str(full)):
        return f"{label}: {plain(amount)} is more than the full marks, {plain(full)}", None
    if amount != amount.quantize(Decimal("0.01")):
        return f"{label}: {value} has more than two decimal places", None
    return None, amount


def _shown(mark, parts):
    if mark is None:
        return ""
    if mark.is_absent:
        return "ABS"
    if parts:
        stored = mark.component_marks or {}
        return " · ".join(plain(stored.get(part.code) or 0) for part in parts)
    return plain(mark.marks_obtained)


def _shown_entry(entry, parts):
    if entry["absent"]:
        return "ABS"
    if parts:
        return " · ".join(plain(entry["parts"][part.code]) for part in parts)
    return plain(entry["score"])


# ------------------------------------------------------------------ saving


def apply(*, user, schedule, section, prepared, students):
    """Save the new and changed rows through the mark grid's own path. Returns what was saved."""
    by_pk = {e.pk: e for e in students}
    rows = []
    for entry in prepared:
        if entry["action"] == "unchanged" or entry["enrollment"] not in by_pk:
            continue
        rows.append(
            (
                by_pk[entry["enrollment"]],
                Decimal(entry["score"]) if entry["score"] not in (None, "") else None,
                entry["absent"],
                entry["version"],
                {code: Decimal(value) for code, value in entry["parts"].items() if value not in (None, "")},
                False,
            )
        )
    if not rows:
        return []
    return save_marks(user=user, schedule=schedule, section=section, rows=rows)
